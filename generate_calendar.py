import sys
import os
import re
import urllib.request
import ssl
from datetime import datetime, time
import openpyxl
from icalendar import Calendar, Event

EXCEL_URL = os.environ.get("EXCEL_URL", "https://www.ur.edu.pl/files/user_directory/307/1%20ROK%20ZIMA%202026-2027.xlsx")

def download_excel(url, output_path):
    """Pobiera plik Excel z podanego adresu URL z pominięciem ew. błędu certyfikatu SSL."""
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, Gecko) Chrome/120.0.0.0 Safari/537.36",
            "Referer": "https://www.ur.edu.pl/"
        }
    )
    with urllib.request.urlopen(req, context=ctx) as response:
        with open(output_path, "wb") as f:
            f.write(response.read())

def parse_schedule(excel_path, target_group="6"):
    """Parsuje rozkład zajęć z pliku Excel dla wskazanej grupy ćwiczeniowej."""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['Arkusz1']

    cell_info = {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell_info[(r, c)] = (ws.cell(r, c).value, r, r, c, c, (r, c))

    for rng in ws.merged_cells.ranges:
        val = ws.cell(rng.min_row, rng.min_col).value
        top_left = (rng.min_row, rng.min_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                cell_info[(r, c)] = (val, rng.min_row, rng.max_row, rng.min_col, rng.max_col, top_left)

    # Znajdowanie bloków dniowych (na podstawie kolumny 2 zawierającej daty)
    day_blocks = []
    current_date = None
    current_start_row = None

    for r in range(1, ws.max_row + 1):
        c2 = ws.cell(r, 2).value
        if isinstance(c2, datetime):
            if current_date is not None:
                day_blocks.append((current_date, current_start_row, r - 1))
            current_date = c2
            current_start_row = r

    if current_date is not None:
        day_blocks.append((current_date, current_start_row, ws.max_row))

    events = []

    for date_val, start_r, end_r in day_blocks:
        g_rows = set()
        for r in range(start_r, end_r + 1):
            c_val = str(ws.cell(r, 3).value or '').strip()
            if not c_val and (r, 3) in cell_info:
                c_val = str(cell_info[(r, 3)][0] or '').strip()
            if c_val == str(target_group):
                g_rows.add(r)
                
        if not g_rows:
            continue
            
        min_g_row = min(g_rows)
        max_g_row = max(g_rows)
        
        seen_top_lefts = set()
        
        for r in range(start_r, end_r + 1):
            for c in range(6, ws.max_column + 1):
                val, min_r, max_r, min_c, max_c, top_left = cell_info[(r, c)]
                if not val:
                    continue
                if top_left in seen_top_lefts:
                    continue
                    
                # Nakładanie się zakresu wierszy z wierszami docelowej grupy
                if not (max_r < min_g_row or min_r > max_g_row):
                    seen_top_lefts.add(top_left)
                    
                    # Kolumna 6 = 07:00, każda kolejna kolumna = 15 minut
                    start_min = 7 * 60 + (min_c - 6) * 15
                    end_min = 7 * 60 + (max_c - 5) * 15
                    
                    t_start = time(start_min // 60, start_min % 60)
                    t_end = time(end_min // 60, end_min % 60)
                    
                    dt_start = datetime.combine(date_val.date(), t_start)
                    dt_end = datetime.combine(date_val.date(), t_end)
                    
                    raw_text = str(val).strip()
                    events.append({
                        "start": dt_start,
                        "end": dt_end,
                        "text": raw_text
                    })
    return events

def generate_ics(events, output_ics_path, title="Plan Zajęć UR - Grupa 6"):
    cal = Calendar()
    cal.add('prodid', '-//Uniwersytet Rzeszowski Plan Zajęć//UR Kalendarz//PL')
    cal.add('version', '2.0')
    cal.add('x-wr-calname', title)
    cal.add('x-wr-timezone', 'Europe/Warsaw')

    for ev in events:
        event = Event()
        raw_text = ev["text"]
        lines = [line.strip() for line in raw_text.split('\n') if line.strip()]
        
        summary = lines[0] if lines else "Zajęcia"
        location = ""
        description = "\n".join(lines)
        
        # Ekstrakcja lokalizacji
        for line in lines:
            if any(kw in line.lower() for kw in ["sala", "ul.", "budynek", "online", "a5", "g4", "a0", "b1", "b2", "c7"]):
                location = line
                break
                
        event.add('summary', summary)
        event.add('dtstart', ev["start"])
        event.add('dtend', ev["end"])
        event.add('description', description)
        if location:
            event.add('location', location)
            
        cal.add_component(event)

    with open(output_ics_path, 'wb') as f:
        f.write(cal.to_ical())
    print(f"Wygenerowano kalendarz ICS: {output_ics_path} ({len(events)} wydarzeń)")

if __name__ == "__main__":
    group = sys.argv[1] if len(sys.argv) > 1 else "6"
    out_path = sys.argv[2] if len(sys.argv) > 2 else f"grupa_{group}.ics"
    
    excel_file = "plan_current.xlsx"
    print(f"Pobieranie aktualnego planu z UR ({EXCEL_URL})...")
    download_excel(EXCEL_URL, excel_file)
    
    print(f"Parsowanie planu dla grupy {group}...")
    events = parse_schedule(excel_file, target_group=group)
    generate_ics(events, out_path, title=f"Plan Zajęć UR - Grupa {group}")
